"""Parse Altium-style pin reports and match core-board FPGA IOs to baseboard functions."""

from __future__ import annotations

import re
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

PIN_HEADERS = ("标号", "pin", "pin number", "designator", "pad")
NAME_HEADERS = ("引脚名称", "pin name", "name", "signal")
NET_HEADERS = ("网络", "net", "net name", "netname")
DIFF_HEADERS = ("差分", "总线", "diff", "bus")
ELEC_HEADERS = ("电气", "electrical", "type")

IO_DIFF_RE = re.compile(
    r"^IO_L(\d+)([PN])(?:_[A-Z0-9]+)*_(\d+)$",
    re.IGNORECASE,
)
IO_SE_RE = re.compile(r"^IO_(\d+)_(\d+)$", re.IGNORECASE)
BASE_DIFF_RE = re.compile(
    r"^BANK(\d+)_(\d+)(?:_(?:MRC|SRC|MRCC|SRCC))?_([PN])$",
    re.IGNORECASE,
)
BASE_SE_RE = re.compile(r"^BANK(\d+)_(\d+)$", re.IGNORECASE)
BANK_TOKEN_RE = re.compile(r"BANK(\d+)", re.IGNORECASE)


def _norm(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _header_index(headers: list[str], aliases: tuple[str, ...], fallback: int) -> int:
    lowered = [h.lower().rstrip(".") for h in headers]
    for alias in aliases:
        alias_l = alias.lower()
        for i, h in enumerate(lowered):
            if h == alias_l or alias_l in h:
                return i
    return fallback if fallback < len(headers) else 0


def _load_sheet_rows(source: Path | bytes, filename: str = "") -> tuple[str, list[dict[str, str]]]:
    if isinstance(source, bytes):
        wb = load_workbook(BytesIO(source), data_only=True, read_only=True)
    else:
        wb = load_workbook(source, data_only=True, read_only=True)

    best_name = wb.sheetnames[0]
    best_rows: list[list[Any]] = []
    for name in wb.sheetnames:
        ws = wb[name]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        nonempty = sum(1 for r in rows if any(_norm(c) for c in r))
        if nonempty > sum(1 for r in best_rows if any(_norm(c) for c in r)):
            best_name, best_rows = name, rows
    wb.close()

    if not best_rows:
        return filename, []

    headers = [_norm(c) for c in best_rows[0]]
    pin_i = _header_index(headers, PIN_HEADERS, 0)
    name_i = _header_index(headers, NAME_HEADERS, 1)
    net_i = _header_index(headers, NET_HEADERS, 2)
    diff_i = _header_index(headers, DIFF_HEADERS, 3)
    elec_i = _header_index(headers, ELEC_HEADERS, 4)

    parsed: list[dict[str, str]] = []
    for raw in best_rows[1:]:
        if not raw or not any(_norm(c) for c in raw):
            continue
        def col(idx: int) -> str:
            return _norm(raw[idx]) if idx < len(raw) else ""

        pin = col(pin_i)
        name = col(name_i)
        net = col(net_i)
        if not pin and not name and not net:
            continue
        parsed.append(
            {
                "pin": pin,
                "name": name,
                "net": net,
                "diff": col(diff_i),
                "electrical": col(elec_i),
            }
        )
    return filename, parsed


def _silk_base(name: str) -> str:
    return name.split("/", 1)[0].strip()


def canonical_from_fpga(io_name: str) -> str | None:
    name = _norm(io_name)
    m = IO_DIFF_RE.match(name)
    if m:
        return f"BANK{m.group(3)}_{m.group(1)}_{m.group(2).upper()}"
    m = IO_SE_RE.match(name)
    if m:
        return f"BANK{m.group(2)}_{m.group(1)}"
    return None


def canonical_from_base(pin_name: str) -> str | None:
    name = _silk_base(pin_name)
    if not name or name.upper() == "GND":
        return None
    m = BASE_DIFF_RE.match(name)
    if m:
        return f"BANK{m.group(1)}_{m.group(2)}_{m.group(3).upper()}"
    m = BASE_SE_RE.match(name)
    if m:
        return f"BANK{m.group(1)}_{m.group(2)}"
    return None


def bank_of(canonical: str | None) -> str:
    if not canonical:
        return ""
    m = BANK_TOKEN_RE.search(canonical)
    return m.group(1) if m else ""


def _is_pl_io(name: str) -> bool:
    return bool(canonical_from_fpga(name))


def _is_ps(name: str) -> bool:
    return _norm(name).upper().startswith("PS_")


def detect_kind(rows: list[dict[str, str]]) -> str:
    core = 0
    base = 0
    for r in rows:
        name = r["name"]
        if _is_pl_io(name) or _is_ps(name):
            core += 1
        silk = _silk_base(name)
        if silk.upper().startswith("BANK") or canonical_from_base(name):
            base += 1
    if core > base:
        return "core"
    if base > core:
        return "base"
    return "unknown"


def _search_blob(*parts: str) -> str:
    return " ".join(p for p in parts if p).lower()


def _sort_key(row: dict[str, Any]) -> tuple:
    bank = row.get("bank") or ""
    try:
        bank_n = int(bank)
    except ValueError:
        bank_n = 10_000
    canon = row.get("canonical") or ""
    m = re.search(r"_(\d+)(?:_([PN]))?$", canon)
    pair = int(m.group(1)) if m else 10_000
    pn = {"P": 0, "N": 1}.get((m.group(2) or "").upper(), 2) if m else 2
    return (bank_n, pair, pn, row.get("package_pin") or "", row.get("connector_pin") or "")


def _row(
    *,
    category: str,
    function: str,
    package_pin: str = "",
    io_name: str = "",
    bank: str = "",
    connector_pin: str = "",
    core_silk: str = "",
    canonical: str = "",
    electrical: str = "",
    extra: str = "",
) -> dict[str, Any]:
    return {
        "category": category,
        "function": function,
        "package_pin": package_pin,
        "io_name": io_name,
        "bank": bank,
        "connector_pin": connector_pin,
        "core_silk": core_silk,
        "canonical": canonical,
        "electrical": electrical,
        "extra": extra,
        "search_text": _search_blob(
            function,
            package_pin,
            io_name,
            bank,
            f"BANK{bank}" if bank else "",
            connector_pin,
            core_silk,
            canonical,
            extra,
        ),
    }


def match_tables(
    core_rows: list[dict[str, str]],
    base_rows: list[dict[str, str]],
) -> dict[str, Any]:
    core_by_key: dict[str, list[dict[str, str]]] = {}
    pl_rows: list[dict[str, str]] = []
    fpga_other: list[dict[str, str]] = []

    for r in core_rows:
        key = canonical_from_fpga(r["name"])
        if key:
            core_by_key.setdefault(key, []).append(r)
            pl_rows.append(r)
        else:
            fpga_other.append(r)

    base_by_key: dict[str, list[dict[str, str]]] = {}
    base_other: list[dict[str, str]] = []
    for r in base_rows:
        key = canonical_from_base(r["name"])
        if key:
            base_by_key.setdefault(key, []).append(r)
        else:
            base_other.append(r)

    matched: list[dict[str, Any]] = []
    used_core: set[int] = set()
    used_base: set[int] = set()

    for key, bases in base_by_key.items():
        cores = core_by_key.get(key, [])
        if not cores:
            for b in bases:
                base_other.append(b)
            continue
        for i, b in enumerate(bases):
            c = cores[min(i, len(cores) - 1)]
            used_core.add(id(c))
            used_base.add(id(b))
            silk = _silk_base(b["name"])
            matched.append(
                _row(
                    category="matched",
                    function=b["net"] or silk,
                    package_pin=c["pin"],
                    io_name=c["name"],
                    bank=bank_of(key),
                    connector_pin=b["pin"],
                    core_silk=silk,
                    canonical=key,
                    electrical=b["electrical"] or c["electrical"],
                    extra=c["diff"],
                )
            )

    unmatched_pl: list[dict[str, Any]] = []
    for r in pl_rows:
        if id(r) in used_core:
            continue
        key = canonical_from_fpga(r["name"]) or ""
        net = r["net"]
        if net and not net.upper().startswith("NET"):
            function = net
        else:
            function = r["diff"] or "核心板自用"
        unmatched_pl.append(
            _row(
                category="unrouted",
                function=function,
                package_pin=r["pin"],
                io_name=r["name"],
                bank=bank_of(key),
                core_silk=key,
                canonical=key,
                electrical=r["electrical"],
                extra=r["diff"],
            )
        )

    folded_base: list[dict[str, Any]] = []
    for r in base_other:
        if id(r) in used_base:
            continue
        silk = _silk_base(r["name"])
        folded_base.append(
            _row(
                category="base_other",
                function=r["net"] or silk or "(空网络)",
                connector_pin=r["pin"],
                core_silk=silk,
                electrical=r["electrical"],
            )
        )

    folded_fpga: list[dict[str, Any]] = []
    for r in fpga_other:
        cat = "ps" if _is_ps(r["name"]) else "fpga_other"
        folded_fpga.append(
            _row(
                category=cat,
                function=r["net"] or r["name"],
                package_pin=r["pin"],
                io_name=r["name"],
                electrical=r["electrical"],
                extra=r["diff"],
            )
        )

    matched.sort(key=_sort_key)
    unmatched_pl.sort(key=_sort_key)
    banks = sorted({r["bank"] for r in matched + unmatched_pl if r["bank"]}, key=lambda x: int(x) if x.isdigit() else x)

    return {
        "matched": matched,
        "unrouted": unmatched_pl,
        "base_other": folded_base,
        "fpga_other": [r for r in folded_fpga if r["category"] != "ps"],
        "ps": [],
        "banks": banks,
        "stats": {
            "matched": len(matched),
            "unrouted": len(unmatched_pl),
            "base_other": len(folded_base),
            "fpga_other": sum(1 for r in folded_fpga if r["category"] != "ps"),
            "ps": sum(1 for r in folded_fpga if r["category"] == "ps"),
            "core_pins": len(core_rows),
            "base_pins": len(base_rows),
        },
    }


def analyze_files(
    file_a: Path | bytes,
    file_b: Path | bytes,
    name_a: str = "",
    name_b: str = "",
) -> dict[str, Any]:
    n1, rows1 = _load_sheet_rows(file_a, name_a)
    n2, rows2 = _load_sheet_rows(file_b, name_b)
    k1, k2 = detect_kind(rows1), detect_kind(rows2)

    if k1 == "core" and k2 != "core":
        core_name, core_rows, base_name, base_rows = n1, rows1, n2, rows2
    elif k2 == "core" and k1 != "core":
        core_name, core_rows, base_name, base_rows = n2, rows2, n1, rows1
    elif k1 == "base" and k2 != "base":
        core_name, core_rows, base_name, base_rows = n2, rows2, n1, rows1
    elif k2 == "base" and k1 != "base":
        core_name, core_rows, base_name, base_rows = n1, rows1, n2, rows2
    else:
        # Fall back to whichever looks more FPGA-like.
        s1 = sum(1 for r in rows1 if _is_pl_io(r["name"]))
        s2 = sum(1 for r in rows2 if _is_pl_io(r["name"]))
        if s1 >= s2:
            core_name, core_rows, base_name, base_rows = n1, rows1, n2, rows2
        else:
            core_name, core_rows, base_name, base_rows = n2, rows2, n1, rows1
        k1 = k2 = "guessed"

    report = match_tables(core_rows, base_rows)
    report["core_file"] = core_name
    report["base_file"] = base_name
    report["detected"] = {
        "file_a": {"name": n1, "kind": detect_kind(rows1), "rows": len(rows1)},
        "file_b": {"name": n2, "kind": detect_kind(rows2), "rows": len(rows2)},
        "core_is": core_name,
        "base_is": base_name,
        "guessed": k1 == "guessed",
    }
    return report
